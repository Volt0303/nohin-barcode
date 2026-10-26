"""Builds the receiving workbook template (入荷読み取り.xlsx).

The VBA in excel/vba/ is imported into it once on Windows and the file is
saved as .xlsm (see excel/セットアップ手順.txt). Sheet code names shScan and
shData are set here so the VBA does not depend on the Japanese sheet names.

Usage: python excel/build_excel.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
DIST = HERE / "dist"
OUT = DIST / "入荷読み取り_template.xlsx"
# Japanese Excel imports VBA source as Shift_JIS (cp932) with CRLF line ends.
VBA_SOURCES = ["modScan.bas", "ThisWorkbook.txt"]

NAVY = "1F3A5F"
FONT = "Meiryo UI"
thin = Side(style="thin", color="B7C3D0")
thick = Side(style="medium", color=NAVY)


def build():
    wb = Workbook()

    scan = wb.active
    scan.title = "読み取り"
    scan.sheet_properties.codeName = "shScan"
    scan.sheet_view.showGridLines = False
    scan.column_dimensions["A"].width = 22
    scan.column_dimensions["B"].width = 60
    scan.column_dimensions["C"].width = 4

    scan["A1"] = "入荷バーコード読み取り"
    scan["A1"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
    scan.row_dimensions[1].height = 34

    scan["A3"] = "読み取り欄"
    scan["A3"].font = Font(name=FONT, size=12, bold=True)
    scan["A3"].alignment = Alignment(vertical="center")
    scan["B3"].font = Font(name=FONT, size=20)
    scan["B3"].fill = PatternFill("solid", fgColor="FFF8D6")
    scan["B3"].border = Border(left=thick, right=thick, top=thick, bottom=thick)
    scan["B3"].alignment = Alignment(vertical="center")
    scan["B3"].number_format = "@"
    scan.row_dimensions[3].height = 40

    # Scanners type like a keyboard; with the Japanese IME on the input gets
    # garbled, so the IME is disabled for the input cell.
    dv = DataValidation(allow_blank=True, imeMode="disabled")
    scan.add_data_validation(dv)
    dv.add("B3")

    scan["A5"] = "結果"
    scan["A5"].font = Font(name=FONT, size=12, bold=True)
    scan["B5"] = "待機中"
    scan["B5"].font = Font(name=FONT, size=12, color="555555")
    scan.row_dimensions[5].height = 26

    scan["A7"] = "本日の読み取り件数"
    scan["A7"].font = Font(name=FONT, size=12, bold=True)
    scan["B7"] = '=COUNTIF(入荷データ!A:A,">="&TODAY())'
    scan["B7"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    scan["B7"].alignment = Alignment(horizontal="left")
    scan["B7"].number_format = '0"件"'

    notes = [
        "【使い方】",
        "1. 黄色の読み取り欄を選択した状態で、納品書のバーコードを読み取ります。",
        "2. 「入荷データ」シートに1行追加され、自動で保存されます。",
        "3. 合格数は納入数が初期値です。検査後、「入荷データ」シートで修正してください。",
        "4. 同じ注番・同じ納入数を再度読み取ると確認メッセージが表示されます。",
    ]
    for i, text in enumerate(notes):
        cell = scan.cell(row=10 + i, column=1, value=text)
        cell.font = Font(name=FONT, size=10, bold=(i == 0), color="444444")

    data = wb.create_sheet("入荷データ")
    data.sheet_properties.codeName = "shData"
    headers = ["読み取り日時", "注文番号", "納入数", "合格数", "読み取りPC"]
    widths = [20, 14, 10, 10, 16]
    for col, (text, width) in enumerate(zip(headers, widths), start=1):
        cell = data.cell(row=1, column=col, value=text)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
        data.column_dimensions[cell.column_letter].width = width
    data.row_dimensions[1].height = 22
    data.freeze_panes = "A2"
    data.auto_filter.ref = "A1:E1"
    # 合格数 is the one column people edit by hand.
    data["D1"].fill = PatternFill("solid", fgColor="B8860B")

    DIST.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT)
    for name in VBA_SOURCES:
        text = (HERE / "vba" / name).read_text(encoding="utf-8")
        dst = DIST / name
        dst.write_bytes(text.replace("\r\n", "\n").replace("\n", "\r\n").encode("cp932"))
        print("saved", dst)


if __name__ == "__main__":
    build()
